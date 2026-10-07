"""Editable Word version of the patient-impact decision checklist handout."""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK, TEAL, DEEP = RGBColor(0x17, 0x26, 0x2B), RGBColor(0x0F, 0x7C, 0x74), RGBColor(0x0E, 0x3B, 0x3D)
rows = [
    ("Need", "Will it help me?", "Is the need validated with patients and clinicians? Is the need statement solution-neutral?", "Need statement; observation notes"),
    ("Evidence", "Is there proof, for people like me?", "Does the evidence match the exact claim, users and device version? Are failure modes controlled? What is demonstrated, reported, planned?", "Protocol; subgroup results; risk file (ISO 14971); version log"),
    ("People", "Can I freely refuse?", "Is consent independent of any dependent relationship, in the person’s language? Are conflicts disclosed and managed?", "Consent form and process; COI management plan; CTRI entry"),
    ("Permission", "Is it allowed for this use?", "Which CDSCO permission covers this use: test licence (MD-13), clinical investigation (MD-23), manufacturing licence?", "Class; licence or permission; registered ethics committee approval"),
    ("Access", "Can I afford and use it?", "What is the total cost of use? Who is excluded by its demands (vision, dexterity, language, power, phone)? Do marketing claims match the evidence?", "Cost of a full course; inclusive-design audit; service plan; claims review"),
    ("Safety", "Who answers later?", "Who records, reports (MvPI) and acts on harm? How are updates validated, released and rolled back?", "Complaint log; MvPI route; change-control and rollback plan"),
]

def shade(cell, hex_):
    tc = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_); tc.append(s)

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Cm(1.3))
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10); st.font.color.rgb = INK

h = doc.add_heading("Patient-impact decision checklist", level=1)
for r in h.runs: r.font.color.rgb = DEEP; r.font.name = "Cambria"
p = doc.add_paragraph("Bridging the Gap: From Innovation to Patient Impact · For every “Not yet”: name the owner and the condition for proceeding.")
p.runs[0].font.color.rgb = TEAL

t = doc.add_table(rows=1, cols=6); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [Cm(2.4), Cm(3.6), Cm(8.2), Cm(5.4), Cm(3.0), Cm(4.4)]
for i, txt in enumerate(["Stage", "Patient’s question", "Ask before proceeding", "Evidence to see", "Owner", "Go?"]):
    c = t.rows[0].cells[i]; c.text = ""; run = c.paragraphs[0].add_run(txt); run.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, "0E3B3D")
for stage, q, ask, ev in rows:
    cells = t.add_row().cells
    vals = [stage, f"“{q}”", ask, ev, "", "☐ Yes\n☐ Not yet — condition:"]
    for i, v in enumerate(vals):
        cells[i].text = v
        if i == 0:
            cells[i].paragraphs[0].runs[0].bold = True; cells[i].paragraphs[0].runs[0].font.color.rgb = TEAL
        if i == 1:
            cells[i].paragraphs[0].runs[0].italic = True
for row in t.rows:
    for i, w in enumerate(widths): row.cells[i].width = w

doc.add_paragraph()
c = doc.add_paragraph(); r = c.add_run("My commitment  "); r.bold = True
c.add_run("Before ____________________, I will ______________________________.  Owner: ______________.  I proceed only if ______________________________.")
k = doc.add_paragraph(); r = k.add_run("Label every claim: "); r.bold = True
k.add_run("Demonstrated · Reported, preliminary · Planned · Not established.   India anchors: MDR-2017 (MD-13; MD-22/23; MD-5/MD-9) · ethics committee registered with CDSCO · CTRI · ICMR 2017 · MvPI 1800-180-3024 · UCMPMD 2024 · DPDP Rules 2025.")
f = doc.add_paragraph("Workshop synthesis adapted from Stanford Biodesign, Harvard MRCT Center, Cambridge Engineering Design Centre, IDEAL-D, Declaration of Helsinki (2024), ICMR (2017) and CDSCO guidance. Not a regulatory checklist or certification. A licence, patent, grant or institutional association is not evidence of readiness. Checked 7 October 2026.")
f.runs[0].font.size = Pt(8); f.runs[0].font.color.rgb = RGBColor(0x6E, 0x7F, 0x7D)
doc.save("../docs/Decision_Checklist_Handout.docx")
print("saved")
