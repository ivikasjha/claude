"""Count visible words per slide (excluding the source footnote and slide number)."""
import re, sys
from pptx import Presentation
p = Presentation(sys.argv[1])
rows = []
for i, sl in enumerate(p.slides, 1):
    words = []
    for sh in sl.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx in (102, 4294967295):
            continue  # source footnote, slide number
        if sh.has_text_frame:
            words += sh.text_frame.text.split()
        if sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    words += c.text.split()
    w = [x for x in words if re.search(r"[A-Za-z₹]", x) and len(x.strip("“”‘’.,?!:;()")) > 1]
    rows.append((i, len(w)))
    print(f"{i:2d} {len(w):3d}", flush=True)
main = [n for i, n in rows if i <= 25]
print("main slides <=30 words:", sum(1 for n in main if n <= 30), "of", len(main))
